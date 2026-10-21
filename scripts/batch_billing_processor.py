import os
import re
import glob
import json
import sys
import openpyxl
import pandas as pd
import gspread
from google.oauth2.service_account import Credentials

def clean_vendor_name(name):
    if not name:
        return "기타"
    cleaned = str(name).strip()
    cleaned = re.sub(r'[\(（]?[주株式]+[\)）]?', '', cleaned)
    cleaned = cleaned.replace('㈜', '').replace('주식회사', '').strip()
    return cleaned

def clean_number(val):
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    val_str = str(val).strip().replace(',', '').replace('₩', '').replace('원', '')
    if val_str.startswith('(') and val_str.endswith(')'):
        val_str = '-' + val_str[1:-1]
    match = re.search(r'[-+]?\d*\.?\d+', val_str)
    return float(match.group()) if match else 0.0

def map_standard_craft(job_title):
    title = str(job_title).strip()
    if '배관' in title:
        return '배관공'
    elif '용접' in title:
        return '용접사'
    elif '도비' in title:
        return '도비공'
    elif '보통' in title:
        return '보통인부'
    elif any(k in title for k in ['관리', '보조', '공사보조', '안전보조', '관리보조']):
        return '관리자'
    return '기타'

def find_sheet_by_keywords(wb, keywords):
    for name in wb.sheetnames:
        norm_name = re.sub(r'\s+', '', name)
        if any(re.sub(r'\s+', '', kw) in norm_name for kw in keywords):
            return wb[name]
    return None

def extract_expense_amount(sheet):
    if not sheet:
        return 0.0
    for cell_pos in ['I2', 'H2', 'J2']:
        try:
            val = clean_number(sheet[cell_pos].value)
            if val > 0:
                return val
        except Exception:
            pass
    for r in range(1, min(sheet.max_row + 1, 50)):
        row_str = " ".join([str(sheet.cell(r, c).value or '') for c in range(1, 15)])
        if any(k in row_str for k in ['합계', '소계', '총액', '지급액']):
            nums = [clean_number(sheet.cell(r, c).value) for c in range(1, 15)]
            valid_nums = [n for n in nums if n > 0]
            if valid_nums:
                return max(valid_nums)
    return 0.0

def extract_quantities(sheet):
    if not sheet:
        return 0.0, 0.0
    try:
        j_val = clean_number(sheet['J50'].value)
        l_val = clean_number(sheet['L50'].value)
        if j_val > 0 or l_val > 0:
            return j_val, l_val
    except Exception:
        pass

    col_now, col_cum = None, None
    for r in range(1, 15):
        row_vals = [str(sheet.cell(r, c).value or '').strip() for c in range(1, 25)]
        for c, val in enumerate(row_vals, 1):
            if '금회' in val and any(k in val for k in ['수량', '물량', 'DIA']):
                col_now = c
            elif '누계' in val and any(k in val for k in ['수량', '물량', 'DIA']):
                col_cum = c
        if col_now and col_cum:
            break

    dia_now, dia_cum = 0.0, 0.0
    for r in range(10, sheet.max_row + 1):
        cell_label = str(sheet.cell(r, 1).value or '') + str(sheet.cell(r, 2).value or '') + str(sheet.cell(r, 3).value or '')
        if '합계' in cell_label or '총계' in cell_label:
            if col_now and col_cum:
                dia_now = clean_number(sheet.cell(r, col_now).value)
                dia_cum = clean_number(sheet.cell(r, col_cum).value)
            else:
                row_nums = [clean_number(sheet.cell(r, c).value) for c in range(4, 20)]
                valid_nums = [n for n in row_nums if n > 0]
                if len(valid_nums) >= 2:
                    dia_now = valid_nums[0]
                    dia_cum = valid_nums[1]
            break
    return dia_now, dia_cum

def parse_single_excel(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    file_name = os.path.basename(file_path)

    ws_cover = find_sheet_by_keywords(wb, ['표지', '용역청구서표지'])
    ws_gap = find_sheet_by_keywords(wb, ['갑지', '용역청구서(갑지)']) or wb.active

    pjt_name = ws_cover['B7'].value if (ws_cover and ws_cover['B7'].value) else "P4 Ph4 GAS 1공구 마감공사"

    raw_vendor = ws_cover['C15'].value if (ws_cover and ws_cover['C15'].value) else ""
    if not raw_vendor:
        for v in ['대성이엔지', '태성하이텍', '리도시스템']:
            if v in file_name:
                raw_vendor = v
                break
    vendor_name = clean_vendor_name(raw_vendor)

    raw_month = ws_cover['C14'].value if (ws_cover and ws_cover['C14'].value) else file_name
    month_match = re.search(r'(\d{2,4})년?\s*(\d{1,2})월?', str(raw_month))
    if month_match:
        y = month_match.group(1)
        y = f"20{y}" if len(y) == 2 else y
        m = str(month_match.group(2)).zfill(2)
        claim_month = f"{y}-{m}"
    else:
        claim_month = "2026-09"

    claim_cost = clean_number(ws_gap['D13'].value)
    if claim_cost == 0.0:
        for r in range(10, 20):
            val = clean_number(ws_gap.cell(r, 4).value)
            if val > 0:
                claim_cost = val
                break

    ws_rent = find_sheet_by_keywords(wb, ['숙소임차', '임차료', '숙소임차료'])
    ws_mgt = find_sheet_by_keywords(wb, ['숙소관리', '관리비', '숙소관리비'])
    ws_meals = find_sheet_by_keywords(wb, ['식대', '식대비용'])

    exp_rent = extract_expense_amount(ws_rent)
    exp_mgt = extract_expense_amount(ws_mgt)
    exp_dorm = exp_rent + exp_mgt
    exp_meals = extract_expense_amount(ws_meals)
    exp_total = exp_dorm + exp_meals

    ws_qty = find_sheet_by_keywords(wb, ['물량내역서', '물량', '내역서'])
    dia_now, dia_cum = extract_quantities(ws_qty)

    craft_data = {
        '배관공': {'공수': 0.0, '금액': 0.0},
        '용접사': {'공수': 0.0, '금액': 0.0},
        '관리자': {'공수': 0.0, '금액': 0.0},
        '보통인부': {'공수': 0.0, '금액': 0.0},
        '도비공': {'공수': 0.0, '금액': 0.0}
    }

    work_sheet = find_sheet_by_keywords(wb, ['근로내역', '노무비명세서', '노무비', '출력일보'])
    if work_sheet:
        header_row = 5
        job_col, md_col, amt_col = 4, 39, 41
        for r in range(1, 10):
            row_vals = [str(work_sheet.cell(r, c).value or '') for c in range(1, 45)]
            if any('직종' in v for v in row_vals) and any('공수' in v for v in row_vals):
                header_row = r
                for c, val in enumerate(row_vals, 1):
                    if '직종' in val: job_col = c
                    elif '공수' in val: md_col = c
                    elif any(k in val for k in ['총액', '지급액', '합계', '금액']): amt_col = c
                break

        for r in range(header_row + 1, work_sheet.max_row + 1):
            name_val = work_sheet.cell(r, 2).value
            job_val = work_sheet.cell(r, job_col).value
            md_val = clean_number(work_sheet.cell(r, md_col).value)
            amt_val = clean_number(work_sheet.cell(r, amt_col).value)

            if not name_val or any(k in str(name_val) for k in ['합계', '소계']):
                continue

            std_job = map_standard_craft(job_val)
            if std_job in craft_data and md_val > 0:
                craft_data[std_job]['공수'] += md_val
                craft_data[std_job]['금액'] += amt_val

    total_md = sum(v['공수'] for v in craft_data.values())
    total_labor = sum(v['금액'] for v in craft_data.values())

    return {
        "파일명": file_name,
        "PJT명": pjt_name,
        "협력업체명": vendor_name,
        "청구월": claim_month,
        "금회_위탁비용": claim_cost,
        "경비_숙소임차료": exp_rent,
        "경비_숙소관리비": exp_mgt,
        "경비_숙소소계": exp_dorm,
        "경비_식대": exp_meals,
        "경비_총합계": exp_total,
        "금회_DIA_수량": dia_now,
        "누계_DIA_수량": dia_cum,
        "배관공_공수": craft_data['배관공']['공수'],
        "배관공_금액": craft_data['배관공']['금액'],
        "용접사_공수": craft_data['용접사']['공수'],
        "용접사_금액": craft_data['용접사']['금액'],
        "관리자_공수": craft_data['관리자']['공수'],
        "관리자_금액": craft_data['관리자']['금액'],
        "보통인부_공수": craft_data['보통인부']['공수'],
        "보통인부_금액": craft_data['보통인부']['금액'],
        "도비공_공수": craft_data['도비공']['공수'],
        "도비공_금액": craft_data['도비공']['금액'],
        "실투입_총공수": total_md,
        "실투입_총노무비": total_labor
    }

def main():
    excel_files = glob.glob('data/*.xlsx') + glob.glob('data/*.xls')
    if not excel_files:
        print("[INFO] data/ 폴더에 처리할 기성 파일이 없습니다.")
        return

    print(f"[INFO] 총 {len(excel_files)}개의 기성 파일을 일괄 처리합니다.")
    records = []
    for f in excel_files:
        try:
            rec = parse_single_excel(f)
            records.append(rec)
            print(f" [성공] {rec['협력업체명']} ({rec['청구월']}): 위탁 {rec['금회_위탁비용']:,}원 | 경비 {rec['경비_총합계']:,}원 | DIA {rec['금회_DIA_수량']} | 공수 {rec['실투입_총공수']} M/D")
        except Exception as e:
            print(f" [오류] {f} 파싱 중 예외 발생: {e}")

    if not records:
        print("[WARN] 유효한 기성 데이터가 없습니다.")
        return

    df = pd.DataFrame(records)
    df = df.sort_values(by=['청구월', '협력업체명'], ascending=[True, True])

    service_account_json = os.environ.get('GCP_SERVICE_ACCOUNT_KEY')
    spreadsheet_id = os.environ.get('SPREADSHEET_ID')

    if not service_account_json or not spreadsheet_id:
        print("[ERROR] GCP_SERVICE_ACCOUNT_KEY 또는 SPREADSHEET_ID 환경변수가 설정되지 않았습니다.")
        df.to_csv("parsed_billing_summary.csv", index=False, encoding='utf-8-sig')
        print("[INFO] parsed_billing_summary.csv 파일로 로컬 저장 완료.")
        sys.exit(1)

    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive"
    ]

    try:
        creds_dict = json.loads(service_account_json.strip())
        credentials = Credentials.from_service_account_info(creds_dict, scopes=scopes)
        gc = gspread.authorize(credentials)
    except Exception as e:
        print(f"[ERROR] 구글 서비스 계정 인증 실패: {e}")
        sys.exit(1)

    try:
        sh = gc.open_by_key(spreadsheet_id.strip())
        ws = sh.get_worksheet(0)
    except Exception as e:
        print(f"[ERROR] 스프레드시트 열기 실패 (ID 및 권한 확인 필요): {e}")
        sys.exit(1)

    try:
        data_to_write = [df.columns.values.tolist()] + df.fillna(0).values.tolist()
        ws.clear()
        ws.update(range_name='A1', values=data_to_write)
        print(f"[SUCCESS] 구글 스프레드시트에 총 {len(df)}개 기성 데이터 덮어쓰기 적재 완료!")
    except Exception as e:
        print(f"[ERROR] 시트 데이터 쓰기(Update) 실패: {e}")
        sys.exit(1)

if __name__ == '__main__':
    main()
